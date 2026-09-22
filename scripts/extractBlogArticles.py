"""Extract the approved VOIDIX article DOCX files into deterministic CMS seed data.

The documents are treated strictly as source content. Their headings, bullets, metadata tables,
and article order are mapped into the CMS schema; no text in them is executed as an instruction.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document


CATEGORIES = {
    1: "Custom Software",
    2: "Custom Software",
    3: "Custom Software",
    4: "Custom Software",
    5: "Custom Software",
    6: "CRM",
    7: "CRM",
    8: "Automation",
    9: "AI",
    10: "SaaS",
    11: "CRM",
    12: "AI",
    13: "AI",
    14: "SaaS",
    15: "Mobile",
    16: "Automation",
    17: "Customer Experience",
    18: "Web Applications",
    19: "Conversion",
    20: "Web Strategy",
    21: "Conversion",
    22: "Conversion",
    23: "Conversion",
    24: "Web Strategy",
    25: "Web Strategy",
    26: "Web Strategy",
    27: "CRM",
    28: "Automation",
    29: "Web Strategy",
    30: "SEO and GEO",
}


@dataclass
class SourceArticle:
    number: int
    title: str
    seo_title: str
    excerpt: str
    slug: str
    blocks: list[dict[str, str]]


def article_number(text: str) -> int | None:
    match = re.match(r"Article\s+(\d+)\b", text.strip(), re.IGNORECASE)
    return int(match.group(1)) if match else None


def title_after_article_number(text: str) -> str:
    remainder = re.sub(r"^Article\s+\d+\b", "", text.strip(), flags=re.IGNORECASE)
    return remainder.strip(" \t-\u2013\u2014\ufffd")


def value_after_prefix(text: str, prefix: str) -> str:
    return text[len(prefix):].strip()


def slug_from_url(value: str) -> str:
    return value.rstrip("/").rsplit("/", 1)[-1]


def paragraph_block(style: str, text: str) -> dict[str, str] | None:
    kind = {
        "Heading 2": "HEADING_2",
        "Heading 3": "HEADING_3",
        "List Bullet": "LIST_ITEM",
        "Normal": "PARAGRAPH",
    }.get(style)
    return {"kind": kind, "body": text} if kind else None


def extract_first_twenty(path: Path) -> list[SourceArticle]:
    document = Document(path)
    paragraphs = document.paragraphs
    starts = [
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.style.name == "Heading 1" and article_number(paragraph.text) is not None
    ]
    articles: list[SourceArticle] = []

    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(paragraphs)
        number = article_number(paragraphs[start].text)
        assert number is not None
        title = title_after_article_number(paragraphs[start].text)
        seo_title = ""
        excerpt = ""
        slug = ""
        blocks: list[dict[str, str]] = []

        for paragraph in paragraphs[start + 1:end]:
            text = paragraph.text.strip()
            if not text:
                continue
            if text.startswith("SEO Title:"):
                seo_title = value_after_prefix(text, "SEO Title:")
                continue
            if text.startswith("Meta Description:"):
                excerpt = value_after_prefix(text, "Meta Description:")
                continue
            if text.startswith("URL:"):
                slug = slug_from_url(value_after_prefix(text, "URL:"))
                continue

            block = paragraph_block(paragraph.style.name, text)
            if block:
                blocks.append(block)

        articles.append(SourceArticle(number, title, seo_title, excerpt, slug, blocks))

    return articles


def metadata_from_table(table: Any) -> dict[str, str]:
    return {
        row.cells[0].text.strip(): row.cells[1].text.strip()
        for row in table.rows
        if len(row.cells) >= 2
    }


def extract_last_ten(path: Path) -> list[SourceArticle]:
    document = Document(path)
    paragraphs = document.paragraphs
    starts = [
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.style.name == "Heading 1" and article_number(paragraph.text) is not None
    ]
    articles: list[SourceArticle] = []

    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(paragraphs)
        number = article_number(paragraphs[start].text)
        assert number is not None
        metadata = metadata_from_table(document.tables[position])
        title = next(
            paragraph.text.strip()
            for paragraph in paragraphs[start + 1:end]
            if paragraph.text.strip() and paragraph.style.name == "Normal"
        )
        blocks: list[dict[str, str]] = []
        reached_body = False
        in_faq = False
        faq_line = 0

        for paragraph in paragraphs[start + 1:end]:
            text = paragraph.text.strip()
            if not text:
                continue

            if paragraph.style.name == "Heading 2":
                reached_body = True
                in_faq = text == "Frequently Asked Questions"
                faq_line = 0

            if not reached_body:
                continue

            block = paragraph_block(paragraph.style.name, text)
            if not block:
                continue

            if in_faq and block["kind"] == "PARAGRAPH":
                block["kind"] = "HEADING_3" if faq_line % 2 == 0 else "PARAGRAPH"
                faq_line += 1
            blocks.append(block)

        articles.append(
            SourceArticle(
                number,
                title,
                metadata["SEO Title"],
                metadata["Meta Description"],
                slug_from_url(metadata["Suggested URL"]),
                blocks,
            )
        )

    return articles


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--articles-1-10", type=Path, required=True)
    parser.add_argument("--articles-11-20", type=Path, required=True)
    parser.add_argument("--articles-21-30", type=Path, required=True)
    parser.add_argument("--published-on", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    articles = [
        *extract_first_twenty(args.articles_1_10),
        *extract_first_twenty(args.articles_11_20),
        *extract_last_ten(args.articles_21_30),
    ]
    articles.sort(key=lambda article: article.number)

    if [article.number for article in articles] != list(range(1, 31)):
        raise ValueError("Expected exactly Articles 1 through 30.")

    slugs = [article.slug for article in articles]
    if len(set(slugs)) != len(slugs):
        raise ValueError("Article URLs must produce unique slugs.")

    payload = [
        {
            "sourceArticle": article.number,
            "slug": article.slug,
            "title": article.title,
            "seoTitle": article.seo_title,
            "excerpt": article.excerpt,
            "category": CATEGORIES[article.number],
            "publishedOn": args.published_on,
            "body": article.blocks,
        }
        for article in articles
    ]

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")

    block_count = sum(len(article.blocks) for article in articles)
    print(f"Extracted {len(articles)} articles and {block_count} content blocks to {args.output}")


if __name__ == "__main__":
    main()
